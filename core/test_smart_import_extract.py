from io import BytesIO
import subprocess
import time
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase

from core.smart_import_extract import SmartImportInputError, extract_document


class SmartImportExtractionTests(SimpleTestCase):
    def test_pasted_text_is_returned_as_a_named_document(self):
        result = extract_document(text="# 世界书\n雾城存在于北境。")

        self.assertEqual(result.filename, "粘贴文本")
        self.assertEqual(result.text, "# 世界书\n雾城存在于北境。")
        self.assertEqual(result.warnings, [])

    def test_rejects_missing_or_ambiguous_input(self):
        with self.assertRaises(SmartImportInputError):
            extract_document()

        with self.assertRaises(SmartImportInputError):
            extract_document(
                text="角色资料",
                upload=SimpleUploadedFile("data.txt", b"file"),
            )

    def test_txt_removes_utf8_bom(self):
        upload = SimpleUploadedFile("资料.txt", "\ufeff第一段\n第二段".encode("utf-8"))

        result = extract_document(upload=upload)

        self.assertEqual(result.text, "第一段\n第二段")

    def test_markdown_keeps_headings_and_body(self):
        upload = SimpleUploadedFile("设定.md", "# 人物\n\n她住在雾城。".encode("utf-8"))

        result = extract_document(upload=upload)

        self.assertEqual(result.text, "# 人物\n\n她住在雾城。")

    def test_docx_extracts_paragraphs_and_table_cells(self):
        from docx import Document

        source = Document()
        source.add_heading("玩家卡", level=1)
        source.add_paragraph("玩家名：小林")
        table = source.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "偏好"
        table.cell(0, 1).text = "探索"
        raw = BytesIO()
        source.save(raw)
        upload = SimpleUploadedFile("玩家卡.docx", raw.getvalue())

        result = extract_document(upload=upload)

        self.assertIn("玩家卡", result.text)
        self.assertIn("玩家名：小林", result.text)
        self.assertIn("偏好", result.text)
        self.assertIn("探索", result.text)

    def test_doc_uses_antiword_with_bounded_output_and_no_shell(self):
        upload = SimpleUploadedFile("人物.doc", b"legacy word")
        process = MagicMock()
        process.stdout = BytesIO("角色：林雾\n".encode())
        process.wait.return_value = 0

        with patch("core.smart_import_extract.subprocess.Popen", return_value=process) as run:
            result = extract_document(upload=upload)

        self.assertEqual(result.text, "角色：林雾\n")
        self.assertEqual(run.call_args.args[0][0], "antiword")
        self.assertEqual(run.call_args.kwargs["stdout"], subprocess.PIPE)
        self.assertEqual(run.call_args.kwargs["stderr"], subprocess.DEVNULL)
        self.assertFalse(run.call_args.kwargs["shell"])

    def test_unknown_extension_is_rejected_before_running_a_parser(self):
        upload = SimpleUploadedFile("资料.rtf", b"some text")

        with patch("core.smart_import_extract.subprocess.Popen") as run:
            with self.assertRaises(SmartImportInputError):
                extract_document(upload=upload)

        run.assert_not_called()

    def test_doc_output_over_limit_is_terminated(self):
        process = MagicMock()
        process.stdout = BytesIO(b"x" * 5)
        process.wait.return_value = 0
        with patch("core.smart_import_extract.MAX_DOC_OUTPUT_BYTES", 4), patch("core.smart_import_extract.subprocess.Popen", return_value=process):
            with self.assertRaises(SmartImportInputError):
                extract_document(upload=SimpleUploadedFile("long.doc", b"legacy word"))
        process.kill.assert_called_once()

    def test_doc_parser_timeout_terminates_the_child_process(self):
        class SlowOutput:
            def read(self, _size):
                time.sleep(0.03)
                return b""

        process = MagicMock()
        process.stdout = SlowOutput()
        process.wait.return_value = 0
        with patch("core.smart_import_extract.ANTIWORD_TIMEOUT_SECONDS", 0.005), patch("core.smart_import_extract.subprocess.Popen", return_value=process):
            with self.assertRaises(SmartImportInputError):
                extract_document(upload=SimpleUploadedFile("slow.doc", b"legacy word"))
        process.kill.assert_called_once()

    def test_oversized_file_is_rejected(self):
        upload = SimpleUploadedFile("large.txt", b"x" * (5 * 1024 * 1024 + 1))

        with self.assertRaises(SmartImportInputError):
            extract_document(upload=upload)

    def test_extracted_text_over_limit_is_rejected(self):
        with self.assertRaises(SmartImportInputError):
            extract_document(text="x" * 100_001)

    def test_corrupt_docx_is_rejected(self):
        upload = SimpleUploadedFile("broken.docx", b"not a zip archive")

        with self.assertRaises(SmartImportInputError):
            extract_document(upload=upload)

    def test_word_temporary_lock_file_has_a_clear_error(self):
        upload = SimpleUploadedFile("~$穿书设定.docx", b"temporary lock")

        with self.assertRaisesRegex(SmartImportInputError, "临时锁文件"):
            extract_document(upload=upload)

    def test_docx_uncompressed_size_is_bounded_before_document_parser_runs(self):
        archive = MagicMock()
        archive.__enter__.return_value.infolist.return_value = [SimpleNamespace(file_size=26 * 1024 * 1024)]
        with patch("core.smart_import_extract.zipfile.ZipFile", return_value=archive), patch("core.smart_import_extract.Document") as document:
            with self.assertRaises(SmartImportInputError):
                extract_document(upload=SimpleUploadedFile("large.docx", b"zip"))
        document.assert_not_called()

    def test_empty_document_is_rejected(self):
        with self.assertRaises(SmartImportInputError):
            extract_document(text=" \n ")

        with self.assertRaises(SmartImportInputError):
            extract_document(upload=SimpleUploadedFile("empty.txt", b""))

    def test_missing_antiword_returns_installation_guidance(self):
        upload = SimpleUploadedFile("人物.doc", b"legacy word")

        with patch(
            "core.smart_import_extract.subprocess.Popen",
            side_effect=FileNotFoundError("antiword"),
        ):
            with self.assertRaisesRegex(SmartImportInputError, "antiword"):
                extract_document(upload=upload)
