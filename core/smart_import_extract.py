from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


MAX_SOURCE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
ANTIWORD_TIMEOUT_SECONDS = 15
ALLOWED_EXTENSIONS = {".txt", ".md", ".docx", ".doc"}


class SmartImportInputError(ValueError):
    """Raised when smart-import input cannot be safely extracted."""


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    text: str
    warnings: list[str]


def _validate_text(text: str) -> str:
    if not isinstance(text, str):
        raise SmartImportInputError("请粘贴文字或选择一个文件")
    if len(text) > MAX_TEXT_CHARS:
        raise SmartImportInputError("提取文字不能超过 100,000 个字符")
    if not text.strip():
        raise SmartImportInputError("文档中没有可读取的文字")
    return text


def _read_upload(upload) -> tuple[str, bytes]:
    name = getattr(upload, "name", "")
    filename = Path(str(name).replace("\\", "/")).name
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise SmartImportInputError("仅支持 TXT、MD、DOCX 和 DOC 文件")
    if not filename or len(filename) > 255:
        raise SmartImportInputError("文件名无效")
    size = getattr(upload, "size", None)
    if isinstance(size, int) and size > MAX_SOURCE_BYTES:
        raise SmartImportInputError("上传文件不能超过 5 MiB")
    try:
        raw = upload.read(MAX_SOURCE_BYTES + 1)
    except (OSError, ValueError, AttributeError) as exc:
        raise SmartImportInputError("无法读取上传文件") from exc
    if not isinstance(raw, bytes):
        raise SmartImportInputError("上传文件内容无效")
    if len(raw) > MAX_SOURCE_BYTES:
        raise SmartImportInputError("上传文件不能超过 5 MiB")
    if not raw:
        raise SmartImportInputError("上传文件为空")
    return filename, raw


def _extract_docx(raw: bytes) -> tuple[str, list[str]]:
    try:
        document = Document(BytesIO(raw))
        blocks = []
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                text = Paragraph(child, document).text.strip()
                if text:
                    blocks.append(text)
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                for row in table.rows:
                    cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
    except (PackageNotFoundError, zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        raise SmartImportInputError("DOCX 文件损坏或格式无法读取") from exc
    warnings = []
    if document.inline_shapes:
        warnings.append("文档包含图片或嵌入图形，图片中的文字未提取")
    return "\n".join(blocks), warnings


def _extract_doc(raw: bytes) -> str:
    try:
        with tempfile.TemporaryDirectory(prefix="neon-smart-import-") as temporary:
            document_path = Path(temporary) / "upload.doc"
            document_path.write_bytes(raw)
            result = subprocess.run(
                ["antiword", str(document_path)],
                cwd=temporary,
                capture_output=True,
                check=False,
                encoding="utf-8",
                errors="replace",
                shell=False,
                timeout=ANTIWORD_TIMEOUT_SECONDS,
            )
    except FileNotFoundError as exc:
        raise SmartImportInputError("服务器缺少 antiword，暂时无法读取 DOC 文件") from exc
    except subprocess.TimeoutExpired as exc:
        raise SmartImportInputError("DOC 文件解析超时，请另存为 DOCX 后重试") from exc
    except OSError as exc:
        raise SmartImportInputError("DOC 文件解析失败，请另存为 DOCX 后重试") from exc
    if result.returncode != 0:
        raise SmartImportInputError("DOC 文件损坏或格式无法读取，请另存为 DOCX 后重试")
    return result.stdout


def extract_document(*, text: str | None = None, upload=None) -> ExtractedDocument:
    if (text is None) == (upload is None):
        raise SmartImportInputError("请粘贴文字或选择一个文件")
    if text is not None:
        return ExtractedDocument("粘贴文本", _validate_text(text), [])

    filename, raw = _read_upload(upload)
    extension = Path(filename).suffix.lower()
    warnings = []
    if extension in {".txt", ".md"}:
        try:
            extracted = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise SmartImportInputError("文本文件需使用 UTF-8 编码") from exc
    elif extension == ".docx":
        extracted, warnings = _extract_docx(raw)
    else:
        extracted = _extract_doc(raw)
    cleaned = _validate_text(extracted)
    return ExtractedDocument(filename, cleaned, warnings)
